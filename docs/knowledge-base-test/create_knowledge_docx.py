from pathlib import Path

from docx import Document
from docx.enum.text import WD_BREAK
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


BASE_DIR = Path(__file__).resolve().parent
MD_PATH = BASE_DIR / "12306智能客服知识库测试数据.md"
DOCX_PATH = BASE_DIR / "12306智能客服知识库测试数据.docx"


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_width(table, widths):
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.autofit = False
    tbl = table._tbl
    tbl_pr = tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:type"), "dxa")
    tbl_w.set(qn("w:w"), str(sum(widths)))
    tbl_ind = OxmlElement("w:tblInd")
    tbl_ind.set(qn("w:w"), "120")
    tbl_ind.set(qn("w:type"), "dxa")
    tbl_pr.append(tbl_ind)
    grid = tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)
    for row in table.rows:
        for idx, cell in enumerate(row.cells):
            cell.width = Inches(widths[idx] / 1440)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell)


def add_table(doc, headers, rows, widths):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    table.allow_autofit = False
    set_table_width(table, widths)
    hdr = table.rows[0].cells
    for idx, header in enumerate(headers):
        hdr[idx].text = header
        set_cell_shading(hdr[idx], "E8EEF5")
        for paragraph in hdr[idx].paragraphs:
            for run in paragraph.runs:
                run.bold = True
    for row in rows:
        cells = table.add_row().cells
        for idx, value in enumerate(row):
            cells[idx].text = value
    for row in table.rows:
        for cell in row.cells:
            set_cell_margins(cell)
    return table


def configure_styles(doc):
    section = doc.sections[0]
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.25

    for name, size, color, before, after in [
        ("Heading 1", 16, "2E74B5", 18, 10),
        ("Heading 2", 13, "2E74B5", 14, 7),
        ("Heading 3", 12, "1F4D78", 10, 5),
    ]:
        style = styles[name]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True


def add_title(doc):
    title = doc.add_paragraph()
    title.paragraph_format.space_after = Pt(4)
    run = title.add_run("12306 智能客服知识库测试数据")
    run.font.name = "Calibri"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    run.font.size = Pt(22)
    run.font.bold = True
    run.font.color.rgb = RGBColor.from_string("0B2545")

    subtitle = doc.add_paragraph()
    subtitle.paragraph_format.space_after = Pt(12)
    run = subtitle.add_run("用于验证上传解析、结构切片、向量化、RAG 检索和智能客服回复效果")
    run.font.color.rgb = RGBColor.from_string("555555")

    note = doc.add_paragraph()
    note.paragraph_format.left_indent = Inches(0.12)
    note.paragraph_format.space_after = Pt(12)
    note_run = note.add_run("说明：本文档为仿 12306 项目的测试资料，内容为业务模拟规则，不代表铁路官方最新规定。")
    note_run.bold = True


def add_bullets(doc, items):
    for item in items:
        doc.add_paragraph(item, style="List Bullet")


def build_docx():
    doc = Document()
    configure_styles(doc)
    add_title(doc)

    doc.add_heading("一、账号与用户体系", level=1)
    doc.add_heading("注册、登录与实名", level=2)
    doc.add_paragraph("用户可以通过手机号、用户名和密码完成注册。注册时应填写真实姓名、证件类型、证件号码和手机号。手机号用于接收登录校验、订单通知和支付结果提醒。")
    add_bullets(doc, [
        "忘记密码时，可以通过绑定手机号完成密码重置。",
        "登录态过期时，需要重新登录。",
        "实名认证失败通常由姓名证件不匹配、证件类型选择错误、证件号码格式错误导致。",
        "客服不得要求用户提供完整身份证号、银行卡号或支付密码。",
    ])

    doc.add_heading("常用联系人与乘车人", level=2)
    doc.add_paragraph("用户可以维护常用乘车人，用于快速下单。删除常用乘车人不会删除历史订单；修改乘车人证件号后，已生成订单不会自动变更。")
    add_bullets(doc, [
        "用户问“怎么给家人买票”时，应提示先添加常用乘车人。",
        "未完成实名认证的乘车人可能无法购票。",
        "智能客服展示乘车人信息时应进行脱敏。",
    ])

    doc.add_heading("二、车票查询知识", level=1)
    doc.add_heading("出发地、到达地与车站", level=2)
    doc.add_paragraph("省市只用于缩小车站选择范围，真正参与余票查询的是具体车站名称。只输入省份通常不能直接查到车次，因为车次售票区间按车站匹配。")
    add_bullets(doc, [
        "选择“江苏省 南京市”后，应展示南京、南京南等车站。",
        "选择“北京市 北京市”后，应展示北京、北京南、北京西、北京朝阳等车站。",
        "用户问“为什么按省查不到票”，应解释需要选择具体出发站和到达站。",
    ])

    add_table(doc, ["类型", "常见前缀", "特点"], [
        ["高铁", "G", "速度快，常见席别有商务座、一等座、二等座"],
        ["动车", "D", "速度较快，常见席别有一等座、二等座"],
        ["城际", "C", "城市群短途运行较多"],
        ["特快", "T", "普速列车，运行时间较长"],
        ["快速", "K", "普速列车，票价相对较低"],
    ], [1800, 1800, 5760])

    doc.add_heading("余票状态", level=2)
    add_bullets(doc, [
        "“有票”表示余票充足或仍可购买。",
        "“少量”表示余票较少，建议尽快下单。",
        "“无票”表示当前席别不可购买。",
        "“候补”表示可提交候补订单等待系统兑现。",
    ])

    doc.add_heading("三、购票业务", level=1)
    doc.add_heading("在线购票流程", level=2)
    for step in [
        "选择出发站、到达站和乘车日期。",
        "查询车次列表，选择车次和席别。",
        "选择乘车人，可选指定座位。",
        "提交订单，并在支付时限内完成支付。",
        "支付成功后，订单状态变为已支付或已出票。",
    ]:
        doc.add_paragraph(step, style="List Number")
    doc.add_heading("选择席位", level=2)
    add_bullets(doc, [
        "指定座位优先尝试锁定指定座位。",
        "指定座位已占用时，不应自动改配其他座位，除非用户重新确认。",
        "未指定座位时，系统按可用座位自动分配。",
        "下单失败时，已扣减的余票和锁定的座位应回滚。",
    ])

    doc.add_heading("候补购票", level=2)
    doc.add_paragraph("候补购票是指当前无票时，用户提交候补需求。当有其他用户退票、改签或系统释放余票时，系统按候补规则尝试兑现。候补不保证一定成功。")

    doc.add_heading("四、支付相关", level=1)
    doc.add_heading("支付渠道与异常", level=2)
    add_table(doc, ["字段", "含义", "客服解释"], [
        ["ALI_PAY", "支付宝支付", "支付宝付款或支付宝扣款问题"],
        ["WECHAT_PAY", "微信支付", "微信付款或微信扣款问题"],
    ], [1800, 2500, 5060])
    add_bullets(doc, [
        "支付成功回调应幂等处理，重复消息不应导致异常。",
        "已支付订单再次收到支付成功消息，应保持已支付。",
        "已关闭订单收到支付成功消息，不应直接改成已支付，应进入人工核查或退款流程。",
        "用户扣款成功但订单未出票，应建议等待系统同步，仍异常时联系人工客服。",
    ])

    doc.add_heading("五、退票规则", level=1)
    doc.add_heading("可退票条件与手续费", level=2)
    doc.add_paragraph("订单必须为已支付或已出票状态，且符合退票时间规则，才可发起退票。未支付订单不需要退票，可直接取消或等待超时关闭。")
    add_table(doc, ["退票时间", "手续费规则"], [
        ["开车前 8 天及以上", "不收退票费"],
        ["开车前 48 小时以上", "按票价 5% 收取"],
        ["开车前 24 小时以上、不足 48 小时", "按票价 10% 收取"],
        ["开车前不足 24 小时", "按票价 20% 收取"],
        ["开车后", "通常不可线上退票，特殊情况按人工规则处理"],
    ], [3600, 5760])
    doc.add_heading("全额退款与部分退款", level=2)
    add_bullets(doc, [
        "全额退款时，可由后端按订单明细计算退款金额。",
        "部分退款必须定位到具体订单明细，例如车厢号、座位号、乘车人或明细 ID。",
        "退款金额应由服务端反查计算，不应信任前端传入的金额。",
        "重复退款消息不应重复释放余票。",
    ])

    doc.add_heading("六、改签规则", level=1)
    doc.add_paragraph("改签用于变更乘车日期、车次、席别或到发站。改签必须基于已支付订单进行，且目标车次有余票。改签失败时原订单通常保持不变。")
    add_bullets(doc, [
        "开车前可尝试改签至其他有票车次。",
        "已办理退票或已全额退款的订单不能改签。",
        "新票价高于原票价时，需要补差价。",
        "新票价低于原票价时，差额按支付渠道退回。",
    ])

    doc.add_heading("七、订单管理与安全边界", level=1)
    add_table(doc, ["状态", "含义", "用户可见解释"], [
        ["待支付", "订单已创建但未支付", "请在支付时限内完成支付"],
        ["已支付", "支付成功", "订单已完成支付，可查看乘车信息"],
        ["部分退款", "部分明细已退款", "部分乘车人或席位已退票"],
        ["全部退款", "全部明细已退款", "该订单已全部退票"],
        ["已关闭", "超时、取消或异常关闭", "订单不可继续支付"],
    ], [1600, 2800, 4960])
    add_bullets(doc, [
        "前端应携带 token。",
        "后端应从登录上下文识别当前用户。",
        "查询订单时必须校验订单归属。",
        "智能客服工具查询订单也必须按当前用户校验归属。",
    ])

    doc.add_heading("八、乘车流程", level=1)
    for step in [
        "到达车站。",
        "使用有效证件进行实名核验。",
        "通过安检。",
        "查看候车室和检票口。",
        "按车站广播或屏幕提示检票进站。",
        "按车厢号和座位号乘车。",
    ]:
        doc.add_paragraph(step, style="List Number")
    doc.add_paragraph("如果身份证遗失、损坏或无法识别，用户应按车站规定办理临时身份证明或使用其他支持的有效证件。")

    doc.add_heading("九、智能客服 FAQ", level=1)
    faq = [
        ("可以提前多久买票？", "购票预售期以系统展示和铁路部门公告为准。建议用户在目标乘车日期进入预售期后尽早查询。"),
        ("为什么按省份查不到车次？", "省市只用于筛选车站，实际查询需要选择具体出发站和到达站。"),
        ("扣款了但没有订单怎么办？", "请确认是否使用当前账号登录，再查看我的订单。若支付成功但订单未同步，建议等待 1-3 分钟后刷新。"),
        ("开车后还能退票吗？", "通常不支持线上退票，特殊情况需按车站或人工规则处理。"),
        ("改签失败会影响原票吗？", "通常不会。改签失败时原订单保持不变。"),
        ("客服能帮我查订单吗？", "智能客服可以在用户登录后查询当前账号下的订单，但不能查询其他账号的订单。"),
    ]
    for question, answer in faq:
        doc.add_paragraph(f"Q：{question}", style="Heading 3")
        doc.add_paragraph(f"A：{answer}")

    doc.add_heading("十、意图分类与同义词", level=1)
    add_table(doc, ["意图", "示例说法", "检索关键词"], [
        ["ticket_search", "查票、看余票", "车票查询、车次、出发站、到达站"],
        ["ticket_purchase", "买票、下单", "购票、乘车人、席别、座位"],
        ["order_query", "查订单、订单在哪", "订单状态、我的订单、订单号"],
        ["payment_issue", "扣款没出票", "支付失败、重复支付、支付回调"],
        ["refund_ticket", "退票、退款", "退票手续费、全额退款、部分退款"],
        ["change_ticket", "改签、换车次", "改签条件、差价、改签失败"],
        ["account_login", "登录不了", "登录态、token、密码、注册"],
        ["real_name_auth", "实名认证失败", "证件校验、乘车人、实名"],
    ], [1900, 2500, 4960])

    doc.add_heading("十一、知识库管理说明", level=1)
    add_bullets(doc, [
        "管理员可以上传 Markdown、Word、PDF 或纯文本文件。",
        "文件上传后会生成入库任务，任务经历解析、切片、向量化和完成等状态。",
        "管理员可以查看切片创建时间、内容、文档 ID 和向量维度。",
        "管理员修改切片内容后，应重新生成该切片向量。",
        "如果知识库检索不到内容，先确认文档入库任务是否完成。",
    ])

    doc.add_heading("十二、测试问题建议", level=1)
    for question in [
        "退票手续费怎么算？",
        "开车后还能退票吗？",
        "支付成功但订单还是待支付怎么办？",
        "为什么我按省份查不到车次？",
        "能不能只退一个乘车人的票？",
        "改签失败原来的票还在吗？",
        "身份证丢了还能坐车吗？",
        "为什么我的订单列表没有订单？",
        "微信支付会不会被识别成支付宝？",
        "管理员怎么确认知识库是否入库成功？",
    ]:
        doc.add_paragraph(question, style="List Number")

    footer = doc.sections[0].footer.paragraphs[0]
    footer.text = "仿 12306 智能客服知识库测试资料"
    footer.alignment = 1

    doc.save(DOCX_PATH)
    print(DOCX_PATH)


if __name__ == "__main__":
    build_docx()
